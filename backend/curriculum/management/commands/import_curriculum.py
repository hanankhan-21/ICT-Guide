from pathlib import Path

from django.core.management.base import BaseCommand
from django.db import transaction
from openpyxl import load_workbook

from curriculum.models import Course, Module, Video, Concept


class Command(BaseCommand):
    help = "Import the ICT University curriculum from the syllabus workbook."

    LEVEL_MAP = {
        "preview / review": "orientation",
        "foundation": "foundation",
        "intermediate": "intermediate",
        "applied core": "applied",
        "applied labs": "applied",
        "advanced apprenticeship": "advanced",
        "advanced reference": "advanced",
        "advanced": "advanced",
    }

    def add_arguments(self, parser):
        parser.add_argument(
            "--reset",
            action="store_true",
            help="Delete existing curriculum data before importing.",
        )

    def handle(self, *args, **options):
        project_dir = Path(__file__).resolve().parents[4]

        workbook_path = (
            project_dir
            / "data"
            / "ICT_University_Video_Syllabus_v1.xlsx"
        )

        if not workbook_path.exists():
            self.stderr.write(
                self.style.ERROR(
                    f"Workbook not found: {workbook_path}"
                )
            )
            return

        self.stdout.write(f"Using workbook: {workbook_path}")

        workbook = load_workbook(
            workbook_path,
            read_only=True,
            data_only=True,
        )

        with transaction.atomic():
            if options["reset"]:
                self.stdout.write("Resetting curriculum data...")

                Video.objects.all().delete()
                Module.objects.all().delete()
                Concept.objects.all().delete()
                Course.objects.all().delete()

            course_count = self.import_courses(workbook)

            module_count, video_count, concept_count = (
                self.import_videos(workbook)
            )

        self.stdout.write("")
        self.stdout.write(
            self.style.SUCCESS("ICT curriculum import complete.")
        )

        self.stdout.write(f"Courses: {course_count}")
        self.stdout.write(f"Modules: {module_count}")
        self.stdout.write(f"Videos: {video_count}")
        self.stdout.write(f"Concepts: {concept_count}")

    def import_courses(self, workbook):
        sheet = workbook["University Syllabus"]

        rows = sheet.iter_rows(
            min_row=2,
            values_only=True,
        )

        count = 0

        for row in rows:
            (
                order,
                stage,
                course_name,
                raw_level,
                purpose,
                source_playlists,
            ) = row

            if not course_name:
                continue

            if "—" in course_name:
                code, title = course_name.split("—", 1)
            else:
                parts = course_name.split(" ", 1)
                code = parts[0]
                title = (
                    parts[1]
                    if len(parts) > 1
                    else course_name
                )

            code = code.strip()
            title = title.strip()

            level_key = str(raw_level or "").strip().lower()

            level = self.LEVEL_MAP.get(
                level_key,
                "foundation",
            )

            Course.objects.update_or_create(
                code=code,
                defaults={
                    "title": title,
                    "description": purpose or "",
                    "level": level,
                    "order": int(order or 0),
                },
            )

            count += 1

        return count

    def import_videos(self, workbook):
        sheet = workbook["Video Watchlist"]

        headers = [
            cell.value
            for cell in sheet[1]
        ]

        column = {
            name: index
            for index, name in enumerate(headers)
        }

        rows = sheet.iter_rows(
            min_row=2,
            values_only=True,
        )

        modules_cache = {}
        module_order_counter = {}

        video_count = 0

        for row in rows:
            course_code = row[column["course_code"]]
            playlist_title = row[column["playlist_title"]]

            if not course_code or not playlist_title:
                continue

            try:
                course = Course.objects.get(
                    code=str(course_code).strip()
                )
            except Course.DoesNotExist:
                self.stdout.write(
                    self.style.WARNING(
                        f"Skipping unknown course: {course_code}"
                    )
                )
                continue

            module_key = (
                course.code,
                str(playlist_title).strip(),
            )

            if module_key not in modules_cache:
                current_order = (
                    module_order_counter.get(
                        course.code,
                        0,
                    )
                    + 1
                )

                module_order_counter[
                    course.code
                ] = current_order

                module, _ = Module.objects.get_or_create(
                    course=course,
                    title=str(playlist_title).strip(),
                    defaults={
                        "order": current_order,
                    },
                )

                modules_cache[module_key] = module

            module = modules_cache[module_key]

            title = row[column["title"]]
            video_id = row[column["video_id"]]
            url = row[column["url"]]

            if not video_id or not url:
                continue

            playlist_order = row[
                column["playlist_order"]
            ]

            video_number = row[
                column["video_number"]
            ]

            transcript_status = (
                row[column["transcript_status"]]
                if "transcript_status" in column
                else ""
            )

            source_depth = (
                row[column["source_depth"]]
                if "source_depth" in column
                else ""
            )

            role = (
                row[column["role"]]
                if "role" in column
                else ""
            )

            prerequisites = (
                row[column["prerequisites"]]
                if "prerequisites" in column
                else ""
            )

            notes_value = (
                row[column["notes"]]
                if "notes" in column
                else ""
            )

            notes_parts = []

            if role:
                notes_parts.append(
                    f"Role: {role}"
                )

            if prerequisites:
                notes_parts.append(
                    f"Prerequisites: {prerequisites}"
                )

            if notes_value:
                notes_parts.append(
                    str(notes_value)
                )

            video, _ = Video.objects.update_or_create(
                video_id=str(video_id).strip(),
                defaults={
                    "module": module,
                    "title": (
                        str(title).strip()
                        if title
                        else str(video_id)
                    ),
                    "youtube_url": str(url).strip(),
                    "playlist_title": str(
                        playlist_title
                    ).strip(),
                    "order": int(
                        video_number
                        or playlist_order
                        or 0
                    ),
                    "transcript_status": (
                        str(transcript_status).strip()
                        if transcript_status
                        else ""
                    ),
                    "evidence_status": (
                        str(source_depth).strip()
                        if source_depth
                        else ""
                    ),
                    "notes": "\n".join(
                        notes_parts
                    ),
                    "is_required": True,
                },
            )

            video.concepts.clear()

            core_concepts = (
                row[column["core_concepts"]]
                if "core_concepts" in column
                else None
            )

            if core_concepts:
                concepts = [
                    item.strip()
                    for item in str(
                        core_concepts
                    ).split(",")
                    if item.strip()
                ]

                for concept_name in concepts:
                    concept, _ = (
                        Concept.objects.get_or_create(
                            name=concept_name
                        )
                    )

                    video.concepts.add(
                        concept
                    )

            video_count += 1

        return (
            Module.objects.count(),
            video_count,
            Concept.objects.count(),
        )